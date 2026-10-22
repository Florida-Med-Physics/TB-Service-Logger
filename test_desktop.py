import json, os, tempfile, unittest
from pathlib import Path
from unittest.mock import patch
from core import Engine, FIELDS
from desktop_worker import DesktopWorker
from report_export import write_report

class NativeTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.reports=self.root/'reports';self.screens=self.root/'screens'
        self.reports.mkdir();self.screens.mkdir()
        self.home=self.root/'home';self.output=self.root/'Export Report.xlsx'
        self.engine=Engine(self.home)
        self.engine.configure({'reports':str(self.reports),'screenshots':str(self.screens),'output':str(self.output)})
    def tearDown(self):
        self.engine.db.close()
        import logging
        logger=logging.getLogger('truebeam.'+str(self.home))
        for handler in logger.handlers[:]:handler.close();logger.removeHandler(handler)
        self.temp.cleanup()
    def test_paths_and_running_survive_restart(self):
        e=Engine(self.home)
        self.assertEqual(e.config['reports'],str(self.reports.resolve()))
        self.assertEqual(e.config['screenshots'],str(self.screens.resolve()))
        self.assertEqual(e.config['output'],str(self.output.resolve()))
        self.assertTrue(e.config['running']);e.db.close()
    def test_pdf_png_import_and_10_min_overwrite(self):
        import pymupdf
        from PIL import Image,ImageDraw
        doc=pymupdf.open();page=doc.new_page()
        page.insert_text((40,60),'Issue: Preventive service\nWork performed: Checked machine\nError code: 420219')
        doc.save(self.reports/'sample.pdf');doc.close()
        image=Image.new('RGB',(1800,550),'navy')
        draw=ImageDraw.Draw(image);draw.rectangle((30,260,700,440),fill=(255,255,210))
        draw.text((40,280),'Id: 420219\nName: COL.SW.MLC.TrajDeviationLeafA',fill='black')
        image.save(self.screens/'sample.png')
        os.utime(self.reports/'sample.pdf',(1,1));os.utime(self.screens/'sample.png',(1,1))
        # OCR is independently tested by the packaged self-test; use deterministic text here.
        with patch('screenshot.extract_screenshot',return_value='Id: 420219\nName: COL.SW.MLC.TrajDeviationLeafA\nDescription: MLC error'):
            self.engine.scan(auto_export=False);self.engine.scan(auto_export=False)
        rows=self.engine.rows();self.assertEqual(len(rows),2)
        self.assertTrue(all(not r['error'] for r in rows))
        self.assertFalse(self.output.exists())
        clock=[100.0];w=DesktopWorker(self.home,clock=lambda:clock[0]);w.engine=self.engine
        w.next_scan=99999;w.next_save=100;w.tick()
        self.assertTrue(self.output.exists())
        first=self.output.stat().st_mtime_ns
        clock[0]=699;w.tick();self.assertEqual(self.output.stat().st_mtime_ns,first)
        clock[0]=700;w.tick();self.assertGreater(self.output.stat().st_mtime_ns,first)
        self.assertEqual(len(list(self.root.glob('Export Report*.xlsx'))),1)
    def test_atomic_failure_keeps_previous_report(self):
        self.output.write_bytes(b'previous-log')
        with patch('report_export.os.replace',side_effect=PermissionError('locked')):
            with self.assertRaises(PermissionError):write_report([],self.output)
        self.assertEqual(self.output.read_bytes(),b'previous-log')
        self.assertEqual(list(self.root.glob('tmp*.xlsx')),[])
    def test_fault_report_columns_status_colors_and_landscape(self):
        from openpyxl import load_workbook
        write_report([dict(kind='screenshots',service_date='2026-10-16T12:00:00',fault_code='420219',work_performed='Reset connection',worker_initials='BK',service_status='Resolved',sources=[]),
                      dict(kind='screenshots',service_date='2026-10-15T12:00:00',fault_code='420220',sources=[])],self.output)
        wb=load_workbook(self.output);ws=wb['Fault Reports']
        columns=[c.value for c in ws[4]]
        self.assertEqual(columns[6:],['Varian Suggested Solution','Current Status','Final Resolution','Initials','Status','Screenshot'])
        self.assertEqual(ws.cell(5,columns.index('Final Resolution')+1).value,'Reset connection')
        resolved,unset=ws.cell(5,columns.index('Status')+1),ws.cell(6,columns.index('Status')+1)
        self.assertEqual((resolved.value,unset.value),('Resolved','Open'))
        self.assertEqual(resolved.fill.fgColor.rgb[-6:],'CDEED6');self.assertEqual(unset.fill.fgColor.rgb[-6:],'FDE2C4')
        self.assertEqual(ws.cell(5,1).font.sz,14)
        for sheet in wb.worksheets:
            self.assertEqual(sheet.page_setup.orientation,'landscape');self.assertEqual(sheet.page_setup.fitToWidth,1)
            self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(wb['Service Reports'].freeze_panes,'A5')
    def test_save_failure_retries_and_missing_folder_keeps_paths(self):
        clock=[600.0];w=DesktopWorker(self.home,clock=lambda:clock[0]);w.engine=self.engine
        w.next_scan=9999;w.next_save=600;self.output.write_bytes(b'original')
        with patch('desktop_worker.write_report',side_effect=PermissionError('locked')):
            w.tick()
        self.assertEqual(self.output.read_bytes(),b'original');self.assertEqual(w.next_retry,630)
        clock[0]=630;w.tick();self.assertEqual(w.next_save,1230)
        self.reports.rmdir();self.engine.scan(auto_export=False)
        self.assertEqual(self.engine.config['reports'],str(self.reports.resolve()))
if __name__=='__main__':unittest.main()
