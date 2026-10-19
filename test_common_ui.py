import hashlib,json,tempfile,threading,unittest
from pathlib import Path
from urllib.request import urlopen
from application_host import SharedService,Server
ROOT=Path(__file__).resolve().parent
class CommonInterfaceTests(unittest.TestCase):
 def test_download_and_desktop_manifest_match(self):
  manifest=json.loads((ROOT/'standalone/ui-manifest.json').read_text())
  self.assertEqual(manifest['ui_sha256'],hashlib.sha256((ROOT/'index.html').read_text(encoding='utf-8').encode()).hexdigest())
  page=(ROOT/'standalone/TrueBeam-Local.html').read_bytes()
  self.assertEqual(manifest['html_sha256'],hashlib.sha256(page).hexdigest())
  text=page.decode()
  self.assertIn('Boolean(window.TRUEBEAM_DESKTOP)',text)
  self.assertIn('if(USE_SHARED)initShared();else initStandalone();',text)
  self.assertIn((ROOT/'ui-policy.js').read_text().strip(),(ROOT/'index.html').read_text())
 def test_host_serves_the_identical_downloadable_html(self):
  with tempfile.TemporaryDirectory() as directory:
   service=SharedService(directory,defaults={'reports':'','screenshots':'','output':'','running':False})
   service.unified_html=True;server=Server(('127.0.0.1',0),service)
   thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
   try:
    with urlopen('http://127.0.0.1:'+str(server.server_port)+'/') as response:body=response.read()
    self.assertEqual(body,(ROOT/'standalone/TrueBeam-Local.html').read_bytes())
   finally:server.shutdown();server.server_close();service.close()
 def test_monitoring_status_does_not_prompt_for_excel_output(self):
  with tempfile.TemporaryDirectory() as directory:
   root=Path(directory);reports=root/'reports';screens=root/'screens';reports.mkdir();screens.mkdir()
   service=SharedService(root/'home',defaults={'reports':str(reports),'screenshots':str(screens),'output':'','running':True})
   try:
    service.engine.message='Choose two folders and an Excel output file to begin.'
    service.tick()
    state=service.state()
    self.assertTrue(state['folders_connected'])
    self.assertIn('Monitoring established folders',state['message'])
    self.assertNotIn('Excel output file',state['message'])
   finally:service.close()
 def test_timestamped_logs_only_for_additions(self):
  with tempfile.TemporaryDirectory() as directory:
   root=Path(directory);logs=root/'logs';logs.mkdir()
   service=SharedService(root/'home',defaults={'reports':'','screenshots':'','output':'','running':False})
   service.engine.config['logs']=str(logs)
   records=[dict(id=1,kind='screenshots',sources=[],issue='Test fault')]
   service.state=lambda:{'records':records}
   try:
    service.save_addition_log();self.assertEqual(len(list(logs.glob('*.xlsx'))),1)
    service.save_addition_log();self.assertEqual(len(list(logs.glob('*.xlsx'))),1)
    records.append(dict(id=2,kind='reports',sources=[],issue='Service work'))
    service.save_addition_log();self.assertEqual(len(list(logs.glob('*.xlsx'))),2)
    records[0]['work_performed']='Updated resolution'
    service.save_addition_log();self.assertEqual(len(list(logs.glob('*.xlsx'))),3)
    service.save_addition_log();self.assertEqual(len(list(logs.glob('*.xlsx'))),3)
    from openpyxl import load_workbook
    book=load_workbook(sorted(logs.glob('*.xlsx'))[-1])
    self.assertEqual(book['Fault Reports'].max_row,5)
    self.assertEqual(book['Service Reports'].max_row,5)
    book.close()
   finally:service.close()
if __name__=='__main__':unittest.main()
