"""Template-style XLSX output for the native application (atomic replacement)."""
import os, tempfile
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SERVICE=[
 ('service_date','Malfunction Start',12),('machine_release','Machine Release',12),
 ('work_order','Work Order #',15),('case_number','Case #',11),('category','Category',18),
 ('subject','Subject / Fault Codes',26),('work_performed','Summary of Work',70),
 ('parts_installed','Parts Installed',32),('engineer','Field Engineer',13),
 ('work_hours','Total Work Hrs',10),('downtime_hours','Agreed Downtime Hrs',11),
 ('service_status','Status',22),('notes','Follow-up / Notes',34),('sources','Source Report',32)]
FAULT=[
 ('service_date','Event Date',20),('fault_code','Error Code',15),
 ('additional_codes','Additional Error Codes',18),('fault_name','Fault Name',32),
 ('severity','Severity',14),('issue','Error Details',70),
 ('suggestions','Varian Suggested Solution',60),('current_status','Current Status',45),('work_performed','Final Resolution',50),
 ('worker_initials','Initials',10),('service_status','Status',12),('sources','Screenshot',32)]
# Fault status is Open until staff mark it Resolved; each value prints with its own color.
STATUS_STYLE={'Open':('FDE2C4','7A3300'),'Resolved':('CDEED6','0F5128')}
def fault_status(record):return 'Resolved' if record.get('service_status')=='Resolved' else 'Open'
def date_value(value):
    if not value:return ''
    for fmt in ('%Y-%m-%dT%H:%M:%S','%Y-%m-%d','%m/%d/%Y %I:%M %p','%m/%d/%Y %H:%M','%m/%d/%Y','%d-%b-%Y %I:%M:%S %p','%d-%b-%Y %I:%M %p','%b/%d/%Y'):
        try:return datetime.strptime(value,fmt)
        except ValueError:pass
    try:return datetime.fromisoformat(value).replace(tzinfo=None)
    except ValueError:return value

def write_report(records,output,*,overwrite=True):
    output=Path(output)
    wb=Workbook();wb.remove(wb.active)
    for kind,title,columns in [('reports','Service Reports',SERVICE),('screenshots','Fault Reports',FAULT)]:
        ws=wb.create_sheet(title);ws.sheet_view.showGridLines=False
        ws.append(['TrueBeam — '+title]);ws['A1'].font=Font(name='Arial',size=20,bold=True)
        ws.append(['Updated '+datetime.now().strftime('%m/%d/%Y %H:%M')+'. Saved from the application database.'])
        ws['A2'].font=Font(name='Arial',size=12,color='595959')
        ws.append([]);ws.append([item[1] for item in columns])
        for cell in ws[4]:
            cell.font=Font(name='Arial',size=14,bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='1F4E78')
            cell.alignment=Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[4].height=42
        rows=[r for r in records if r['kind']==kind]
        rows.sort(key=lambda r:str(date_value(r.get('service_date',''))),reverse=True)
        for record in rows:
            values=[]
            for key,_,_ in columns:
                value=record.get(key,'')
                if key=='subject':value='\n'.join(str(record.get(k,'')) for k in ('issue','fault_code','additional_codes') if record.get(k))
                elif key=='sources':value='\n'.join(source['path'] for source in record['sources'])
                elif key=='service_status' and kind=='screenshots':value=fault_status(record)
                elif key in ('service_date','machine_release'):value=date_value(value)
                elif key in ('work_hours','downtime_hours') and value:
                    try:value=float(value)
                    except ValueError:pass
                values.append(value)
            ws.append(values)
            for col,cell in enumerate(ws[ws.max_row]):
                if isinstance(cell.value,str):cell.data_type='s'
                cell.font=Font(name='Arial',size=14)
                cell.alignment=Alignment(wrap_text=True,vertical='top')
                cell.border=Border(bottom=Side(style='thin',color='D9E2F3'))
                if isinstance(cell.value,datetime):cell.number_format='mm/dd/yyyy hh:mm'
                if kind=='screenshots' and columns[col][0]=='service_status':
                    fill,color=STATUS_STYLE[cell.value]
                    cell.fill=PatternFill('solid',fgColor=fill);cell.font=Font(name='Arial',size=14,bold=True,color=color)
                    cell.alignment=Alignment(horizontal='center',vertical='top')
            # Height accounts for wrapped descriptions; full text remains in cells.
            ws.row_dimensions[ws.max_row].height=min(409,max(80,max((len(str(v))//42+1)*20 for v in values)))
            if record['sources']:
                link=ws.cell(ws.max_row,len(columns))
                link.hyperlink=Path(record['sources'][0]['path']).absolute().as_uri()
                link.font=Font(name='Arial',size=14,color='0563C1',underline='single')
        ws.freeze_panes='A5'
        ws.auto_filter.ref='A4:'+get_column_letter(len(columns))+str(max(4,ws.max_row))
        for i,(_,_,width) in enumerate(columns,1):ws.column_dimensions[get_column_letter(i)].width=width
        # Print landscape, every column on one page width, header row repeated on each page.
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_LETTER
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.print_title_rows='4:4';ws.print_options.horizontalCentered=True
        ws.page_margins.left=ws.page_margins.right=0.3;ws.page_margins.top=ws.page_margins.bottom=0.5
    fd,tmp=tempfile.mkstemp(suffix='.xlsx',dir=output.parent);os.close(fd)
    try:
        wb.save(tmp)
        if overwrite:os.replace(tmp,output)
        else:
            import shutil
            with open(tmp,'rb') as source,open(output,'xb') as target:shutil.copyfileobj(source,target)
    finally:
        if os.path.exists(tmp):os.unlink(tmp)
