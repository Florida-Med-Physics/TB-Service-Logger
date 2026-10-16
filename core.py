"""Local ingestion and atomic Excel export. No cloud services required."""
import hashlib, io, json, os, re, sqlite3, tempfile, time
from datetime import datetime
from pathlib import Path

SUPPORTED = {'.png','.jpg','.jpeg','.tif','.tiff','.bmp','.pdf','.docx','.txt'}
EXTRACTOR_VERSION = '3-local-browser-service-reports'

FIELDS = ['service_date','machine','fault_code','additional_codes','fault_name','severity','impact','issue','cause','suggestions','work_performed','service_status','notes','work_order','case_number','machine_release','category','parts_installed','engineer','work_hours','downtime_hours','document_type','worker_initials']

def extract(path):
    suffix = path.suffix.lower()
    if suffix == '.txt':
        return path.read_text(encoding='utf-8', errors='replace')
    if suffix == '.docx':
        from docx import Document
        doc = Document(path)
        return '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
    if suffix == '.pdf':
        import fitz
        with fitz.open(path) as doc:
            pages = []
            for page in doc:
                text = page.get_text(sort=True).strip()
                if not text:
                    from PIL import Image
                    pix = page.get_pixmap(matrix=fitz.Matrix(2,2))
                    text = ocr(Image.open(io.BytesIO(pix.tobytes('png'))))
                pages.append(text)
            return '\n'.join(pages)
    from PIL import Image, ImageSequence
    from screenshot import extract_screenshot
    with Image.open(path) as image:
        return '\n'.join(extract_screenshot(frame.copy()) for frame in ImageSequence.Iterator(image))

def ocr(image):
    from screenshot import recognize
    return recognize(image)

def parse(text):
    if 'Work Order Information' in text or 'CHECKLIST' in text[:1000].upper() or 'PARAMETER' in text[:1000].upper():
        from service_parser import parse_service
        service=parse_service(text)
    else:
        service={}
    result = dict.fromkeys(FIELDS, '')
    labels = {'service_date': r'(?:service date|date)', 'machine': r'(?:machine|system|serial number)', 'issue': r'(?:issue|problem|fault|error)', 'work_performed': r'(?:work performed|corrective action|resolution|action taken)', 'service_status': r'(?:status)'}
    for field, label in labels.items():
        match = re.search(r'^\s*' + label + r'\s*[:=]\s*(.+)$', text, re.I | re.M)
        if match:
            result[field] = match.group(1).strip()[:4000]
    aliases = {'fault_name':'name', 'severity':'severity', 'impact':'impact', 'issue':'description', 'cause':'causes?', 'suggestions':'suggestions?', 'service_date':r'date\s*time'}
    for field, label in aliases.items():
        match = re.search(r'^\s*'+label+r'\s*[:=]\s*(.*?)(?=^\s*(?:id|name|severity|impact|description|causes?|suggestions?|date\s*time)\s*[:=]|\Z)',text,re.I|re.M|re.S)
        if match: result[field]=match.group(1).strip()[:4000]
    from fault_codes import find_codes
    result['fault_code'], codes = find_codes(text)
    result['additional_codes'] = ', '.join(code for code in codes if code != result['fault_code'])
    if not result['issue']:
        match=re.search(r'^\s*CAUSE\s*[:=]\s*(.+)',text,re.I|re.M)
        if match: result['issue']=match.group(1).strip()
    result.update({k:v for k,v in service.items() if v is not None and v != ''})
    return result

class Engine:
    def __init__(self, home):
        self.home = Path(home)
        self.home.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(self.home / 'service-log.sqlite3', check_same_thread=False, timeout=30)
        self.db.row_factory = sqlite3.Row
        self.db.execute('PRAGMA journal_mode=WAL')
        self.db.execute('PRAGMA busy_timeout=30000')
        self.db.execute('CREATE TABLE IF NOT EXISTS records (id INTEGER PRIMARY KEY, path TEXT UNIQUE, kind TEXT, digest TEXT, imported TEXT, text TEXT, error TEXT, review TEXT, service_date TEXT, machine TEXT, fault_code TEXT, fault_name TEXT, severity TEXT, impact TEXT, issue TEXT, cause TEXT, suggestions TEXT, work_performed TEXT, service_status TEXT, notes TEXT)')
        columns = {r[1] for r in self.db.execute('PRAGMA table_info(records)')}
        for column in tuple(FIELDS)+('extractor_version',):
            if column not in columns:
                self.db.execute(f"ALTER TABLE records ADD COLUMN {column} TEXT DEFAULT ''")
        self.db.execute("UPDATE records SET review='Closure documented' WHERE review='Needs review'")
        self.db.execute("UPDATE records SET service_status='Closure documented' WHERE service_status='Needs review'")
        self.db.commit()
        self.db.execute('CREATE TABLE IF NOT EXISTS links (source_id INTEGER PRIMARY KEY, target_id INTEGER NOT NULL)')
        self.config_path = self.home / 'settings.json'
        self.config = json.loads(self.config_path.read_text()) if self.config_path.exists() else {'screenshots':'','reports':'','output':str(self.home / 'TrueBeam Service Log.xlsx'),'running':False}
        self.seen = {}
        self.message = 'Choose two folders and an Excel output file to begin.'
        self.dirty = True

    def configure(self, values):
        config = {k:str(values.get(k,'')).strip() for k in ('screenshots','reports','output')}
        for name in ('screenshots','reports'):
            folder = Path(config[name]).expanduser().resolve()
            if not config[name] or not folder.is_dir():
                raise ValueError('Choose an existing '+name+' folder.')
            config[name] = str(folder)
        screen, report=Path(config['screenshots']),Path(config['reports'])
        if screen == report or screen in report.parents or report in screen.parents:
            raise ValueError('Choose separate folders that do not contain one another.')
        output = Path(config['output']).expanduser().resolve()
        if output.suffix.lower() != '.xlsx' or not output.parent.is_dir():
            raise ValueError('Excel output must end in .xlsx and its parent folder must exist.')
        # Never silently replace a pre-existing user workbook on initial setup.
        if output.exists() and (str(output) != self.config['output'] or not self.config_path.exists()):
            raise ValueError('That workbook already exists. Choose a new filename to preserve it.')
        config['output'] = str(output)
        config['running'] = True
        self.config = config
        self.config_path.write_text(json.dumps(config, indent=2))
        self.dirty = True
        self.seen.clear()

    def rows(self):
        return [dict(r) for r in self.db.execute('SELECT * FROM records ORDER BY id DESC')]

    def events(self):
        groups = {}
        links = dict(self.db.execute('SELECT source_id,target_id FROM links').fetchall())
        rows = self.rows()
        for row in reversed(rows):
            # Only merge automatically when an explicit event timestamp is present.
            normalized = lambda value: re.sub(r'\s+', '', value or '').lower()
            timestamp = row['service_date']
            key = ('fault', normalized(row['machine']), normalized(row['fault_code']), normalized(row['fault_name']), normalized(timestamp), normalized(row['issue'])) if row['fault_code'] and timestamp and row['kind']=='screenshots' else ('file',row['kind'],row['digest'])
            if row['id'] in links: key = ('manual',links[row['id']],row['kind'])
            elif row['id'] in links.values(): key = ('manual',row['id'],row['kind'])
            if key not in groups:
                groups[key] = dict(row, sources=[], possible_duplicate=False)
            if key[0]=='manual' and row['id']==key[1]:
                sources=groups[key]['sources']
                groups[key]=dict(row,sources=sources,possible_duplicate=False)
            groups[key]['sources'].append({'id':row['id'],'path':row['path'],'error':row['error']})
        events=list(groups.values())
        for event in events:
            event['possible_duplicate'] = any(other['id']!=event['id'] and event['fault_code'] and event['fault_code']==other['fault_code'] and event['fault_name']==other['fault_name'] and event['machine']==other['machine'] for other in events)
        return list(reversed(events))

    def merge(self, source, target):
        events = {r['id']:r for r in self.events()}
        if source == target or source not in events or target not in events: raise ValueError('Choose two different existing events.')
        if events[source]['kind'] != events[target]['kind']:
            raise ValueError('Keep reports and print screens separate; duplicates can only be linked within the same tab.')
        for item in events[source]['sources']+events[target]['sources']:
            self.db.execute('INSERT OR REPLACE INTO links VALUES (?,?)',(item['id'],target))
        self.db.commit()
        self.dirty=True

    def update(self, data):
        record_id = int(data['id'])
        review = data.get('review','Closure documented')
        if review not in ('Closure documented','Reviewed'):
            raise ValueError('Invalid review status')
        self.db.execute('UPDATE records SET '+','.join(f'{k}=?' for k in FIELDS)+',review=? WHERE id=?', [str(data.get(k,'')) for k in FIELDS]+[review,record_id])
        self.db.commit()
        self.dirty = True

    def scan(self):
        if not self.config.get('running'):
            return
        warnings = []
        for key in ('screenshots','reports'):
            folder = Path(self.config[key])
            if not folder.is_dir():
                warnings.append(key+' folder is unavailable')
                continue
            def walk_error(error):
                warnings.append(str(error))
            for root, dirs, files in os.walk(folder, onerror=walk_error):
                dirs[:] = [d for d in dirs if not d.startswith('.')]
                for name in files:
                    path = Path(root) / name
                    if path.suffix.lower() not in SUPPORTED or name.startswith(('.', '~$')) or path.is_symlink():
                        continue
                    try:
                        if callable(getattr(self,'progress',None)):
                            self.progress(path)
                        stat = path.stat()
                        signature = (stat.st_size, stat.st_mtime_ns)
                        # Two unchanged scans plus a quiet interval avoid partially copied files.
                        if self.seen.get(str(path)) != signature or time.time()-stat.st_mtime < 2:
                            self.seen[str(path)] = signature
                            continue
                        old = self.db.execute('SELECT * FROM records WHERE path=?', (str(path),)).fetchone()
                        digest = hashlib.sha256(path.read_bytes()).hexdigest()
                        if old and old['digest'] == digest and old['extractor_version'] == EXTRACTOR_VERSION:
                            continue
                        text, error = '', ''
                        try:
                            text = extract(path).strip()
                            if not text:
                                raise ValueError('No readable text found. Inspect the source and retry.')
                        except Exception as exc:
                            error = str(exc)
                        if (path.stat().st_size,path.stat().st_mtime_ns) != signature:
                            continue
                        fields = parse(text)
                        # Preserve edited fields and notes when the source changes; require a fresh review.
                        if old:
                            previous = parse(old['text'] or '')
                            fields = {k:old[k] if old[k] and (old['review']=='Reviewed' or old[k] != previous.get(k,'')) else fields[k] for k in FIELDS}
                        columns = ['path','kind','digest','imported','text','error','review','extractor_version']+FIELDS
                        values = [str(path),key,digest,datetime.now().isoformat(timespec='seconds'),text,error,'Closure documented',EXTRACTOR_VERSION]+[fields[k] for k in FIELDS]
                        self.db.execute('INSERT INTO records ('+','.join(columns)+') VALUES ('+','.join('?' for _ in columns)+') ON CONFLICT(path) DO UPDATE SET '+','.join(k+'=excluded.'+k for k in columns if k!='path'), values)
                        self.db.commit()
                        self.dirty = True
                    except (OSError, sqlite3.Error) as exc:
                        warnings.append(name+': '+str(exc))
        if self.dirty:
            try:
                self.export()
                self.dirty = False
                self.message = 'Excel updated at '+datetime.now().strftime('%H:%M:%S')
            except Exception as exc:
                self.message = 'Entries saved locally; Excel export pending: '+str(exc)
        if warnings:
            self.message = '; '.join(warnings)

    def export(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.table import Table, TableStyleInfo
        wb = Workbook()
        wb.remove(wb.active)
        headers = ['ID']+[('Error code' if k=='fault_code' else 'Additional error codes' if k=='additional_codes' else k.replace('_',' ').title()) for k in FIELDS]+['Review','Possible duplicate','Source count','Source files','Imported','Extraction error']
        events = self.events()
        from openpyxl.utils import get_column_letter
        for kind, title, table_name in [('reports','Service Reports','ServiceReports'),('screenshots','Fault Reports','FaultReports')]:
            ws = wb.create_sheet(title)
            ws.sheet_view.showGridLines = False
            ws.append(headers)
            records = [record for record in events if record['kind']==kind]
            for r in records:
                errors = [item['error'] for item in r['sources'] if item['error']]
                if kind=='screenshots' and not r['fault_code']:
                    errors.append('No unambiguous primary error code found; inspect the original screenshot.')
                values = [r['id']]+[r[k] for k in FIELDS]+[r['review'],'Check' if r['possible_duplicate'] else '',len(r['sources']),'\n'.join(item['path'] for item in r['sources']),r['imported'],'\n'.join(errors)]
                ws.append(values)
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value,str): cell.data_type = 's'
                    cell.alignment = Alignment(vertical='top',wrap_text=True)
                source_cell=ws.cell(ws.max_row,headers.index('Source files')+1)
                source_cell.hyperlink = Path(r['path']).as_uri()
                source_cell.style = 'Hyperlink'
                source_cell.alignment=Alignment(vertical='top',wrap_text=True)
                ws.row_dimensions[ws.max_row].height = 75
                if errors or r['review'] == 'Closure documented':
                    ws.cell(ws.max_row,headers.index('Review')+1).fill = PatternFill('solid',fgColor='FFF0CB')
                if not r['fault_code'] and kind=='screenshots':
                    ws.cell(ws.max_row,headers.index('Error code')+1).fill = PatternFill('solid',fgColor='FFDCDC')
            ws.freeze_panes = 'F2'
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = PatternFill('solid',fgColor='16364A')
                cell.font = Font(color='FFFFFF',bold=True)
                cell.alignment = Alignment(wrap_text=True,vertical='center')
            ws.row_dimensions[1].height = 36
            for index, header in enumerate(headers,1):
                ws.column_dimensions[get_column_letter(index)].width = 48 if header in ('Issue','Cause','Suggestions','Work Performed','Notes','Source files','Extraction error') else 24
            ws.column_dimensions['A'].width=9
            ws.column_dimensions['D'].width=18
            if records:
                table = Table(displayName=table_name,ref=ws.dimensions)
                table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
                ws.add_table(table)
        output = Path(self.config['output'])
        fd, tmp = tempfile.mkstemp(suffix='.xlsx',dir=output.parent)
        os.close(fd)
        try:
            wb.save(tmp)
            os.replace(tmp,output)
        finally:
            if os.path.exists(tmp): os.unlink(tmp)
