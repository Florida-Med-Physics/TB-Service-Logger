import os, tempfile, unittest
from pathlib import Path
from unittest.mock import patch
from core import Engine, parse, extract
from openpyxl import load_workbook

class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        for folder in ['screens','reports','data']: (self.root/folder).mkdir()
        self.engine=Engine(self.root/'data')
        self.engine.configure({'screenshots':str(self.root/'screens'),'reports':str(self.root/'reports'),'output':str(self.root/'log.xlsx')})
    def tearDown(self):self.engine.db.close();self.tmp.cleanup()
    def add(self,name,text,folder='screens'):
        path=self.root/folder/name;path.write_text(text);os.utime(path,(1,1));return path
    def scan(self):self.engine.scan();self.engine.scan()
    def test_import_repeat_edit_export_and_restart(self):
        p=self.add('report.txt','Service date: 2026-10-08\nMachine: TB1\nIssue: test fault\nWork performed: restarted','reports')
        self.scan();self.scan();self.assertEqual(len(self.engine.rows()),1)
        row=self.engine.rows()[0];row['notes']='=1+1';row['review']='Reviewed';self.engine.update(row);self.engine.scan()
        wb=load_workbook(self.root/'log.xlsx');ws=wb['Service Reports']
        self.assertEqual(ws.max_row,2);self.assertEqual(ws.cell(2,14).value,'=1+1');self.assertEqual(ws.cell(2,14).data_type,'s')
        p.write_text(p.read_text()+'\nStatus: complete');os.utime(p,(2,2));self.scan()
        self.assertEqual(self.engine.rows()[0]['notes'],'=1+1');self.assertEqual(self.engine.rows()[0]['review'],'Needs review')
        second=Engine(self.root/'data');self.assertEqual(len(second.rows()),1);second.db.close()
    def test_exact_and_timed_duplicates_and_recurrence(self):
        text='Id: 420219\nName: COL.SW.MLC.TrajDeviationLeafA\nDescription: mismatch\nDateTime: 09-Oct-2026 08:10:40 AM'
        self.add('one.txt',text);self.add('copy.txt',text);self.add('different-ocr.txt',text+'\n')
        self.add('later.txt',text.replace('08:10:40','09:11:00'))
        self.scan();events=self.engine.events()
        self.assertEqual(len(events),2);self.assertEqual(sorted(len(e['sources']) for e in events),[1,3])
        self.assertTrue(all(e['possible_duplicate'] for e in events))
        self.engine.merge(events[0]['id'],events[1]['id']);self.assertEqual(len(self.engine.events()),1)
        self.assertEqual(len(self.engine.events()[0]['sources']),4)
    def test_without_time_only_flag(self):
        self.add('one.txt','ID: 4016\nName: CTRL_OPEN_PATIENT_Q\nCAUSE: selecting queue')
        self.add('two.txt','ID: 4016\nName: CTRL_OPEN_PATIENT_Q\nCAUSE: selecting queue\n')
        self.scan();self.assertEqual(len(self.engine.events()),2);self.assertTrue(self.engine.events()[0]['possible_duplicate'])
    def test_locked_export_retries(self):
        self.add('one.txt','Issue: sample','reports')
        with patch.object(self.engine,'export',side_effect=PermissionError('locked')):
            self.scan();self.assertTrue(self.engine.dirty);self.assertIn('pending',self.engine.message)
        self.engine.scan();self.assertFalse(self.engine.dirty);self.assertEqual(len(self.engine.rows()),1)
    def test_failed_extraction_visible_and_retry(self):
        self.add('bad.png','invalid image')
        self.scan();self.assertTrue(self.engine.rows()[0]['error']);self.assertEqual(len(self.engine.events()),1)
    def test_explicit_fault_fields(self):
        r=parse('Id: 510102\nName: XI.SW.TRANSFER.Error\nSeverity: Minor\nImpact: BeamOn, KV\nDescription: transfer failed\nCauses: network\nSuggestions: restart\nDateTime: 12-Oct-2026 07:06:23 AM')
        self.assertEqual(r['fault_code'],'510102');self.assertEqual(r['cause'],'network');self.assertEqual(r['work_performed'],'');self.assertEqual(r['suggestions'],'restart')
    def test_code_labels_and_ocr_corrections(self):
        cases = [('Error code: 42O2I9','420219'),('Fault ID = 510102','510102'),('Interlock code # 136011','136011'),('1d: 4O16\nNAME: CTRL_OPEN_PATIENT_Q','4016'),('Error code: 4 2 0 2 1 9','420219')]
        for text,expected in cases:
            with self.subTest(text=text): self.assertEqual(parse(text)['fault_code'],expected)
        for text in ['Patient ID: 17625146','Report ID: 4016','Date: 2026-10-16','Description: leaf 34 position -3.07 cm']:
            with self.subTest(text=text): self.assertEqual(parse(text)['fault_code'],'')
    def test_primary_and_secondary_codes(self):
        row=parse('ID: 420219\nName: COL.SW.MLC.TrajDeviationLeafA\nAdditional fault codes: 420219, 100000')
        self.assertEqual(row['fault_code'],'420219'); self.assertEqual(row['additional_codes'],'100000')
        row=parse('Additional fault codes: 420219, 100000')
        self.assertEqual(row['fault_code'],''); self.assertEqual(row['additional_codes'],'420219, 100000')
    def test_two_tabs_and_source_separation(self):
        text='Error code: 420219'
        self.add('screen.txt',text); self.add('report.txt',text,'reports');self.scan()
        self.assertEqual(len(self.engine.events()),2)
        wb=load_workbook(self.root/'log.xlsx')
        self.assertEqual(wb.sheetnames,['Service Reports','Fault Reports'])
        for title in wb.sheetnames:
            ws=wb[title]; self.assertEqual(ws.max_row,2)
            self.assertEqual(ws.cell(2,4).value,'420219')
            self.assertEqual(len(ws.tables),1)
        rows=self.engine.rows()
        with self.assertRaises(ValueError):self.engine.merge(rows[0]['id'],rows[1]['id'])
    def test_extractor_upgrade_reimports_preserving_manual_notes(self):
        self.add('screen.txt','Error code: 420219');self.scan()
        row=self.engine.rows()[0]
        self.engine.db.execute("UPDATE records SET extractor_version='',fault_code='',notes='Manual note' WHERE id=?",(row['id'],)); self.engine.db.commit()
        self.scan();row=self.engine.rows()[0]
        self.assertEqual(row['fault_code'],'420219');self.assertEqual(row['notes'],'Manual note')
    def test_pdf_and_word_extraction(self):
        import fitz
        from docx import Document
        pdf=self.root/'report.pdf'
        doc=fitz.open();page=doc.new_page();page.insert_text((72,72),'Issue: synthetic PDF fault');doc.save(pdf);doc.close()
        self.assertIn('synthetic PDF fault',extract(pdf))
        word=self.root/'report.docx';doc=Document();doc.add_paragraph('Issue: synthetic Word fault');doc.save(word)
        self.assertIn('synthetic Word fault',extract(word))
    def test_preserve_existing_workbook(self):
        other=self.root/'existing.xlsx';other.write_text('keep')
        with self.assertRaises(ValueError):self.engine.configure(dict(self.engine.config,output=str(other)))
        self.assertEqual(other.read_text(),'keep')

if __name__=='__main__':unittest.main()
