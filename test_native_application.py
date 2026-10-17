import base64,csv,json,os,tempfile,time,unittest,zipfile
from pathlib import Path
from unittest.mock import patch
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from PySide6.QtCore import QTimer,QSettings,Qt
from PySide6.QtWidgets import QApplication,QDialog,QDialogButtonBox,QLineEdit,QPlainTextEdit
from native_application import Backend,Api,Window,FAULT_COLUMNS,REPORT_COLUMNS,cell_value
from live_report import create_live_workbook

APP=QApplication.instance() or QApplication([])
def wait_for(predicate,timeout=12):
    deadline=time.monotonic()+timeout
    while time.monotonic()<deadline:
        APP.processEvents();time.sleep(.015)
        if predicate():return
    raise AssertionError('Timed out waiting for native UI')
class NativeTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        for name in ('reports','screens'):(self.root/name).mkdir()
        self.backend=Backend(local=True,home=self.root/'home',defaults={'reports':str(self.root/'reports'),'screenshots':str(self.root/'screens'),'output':str(self.root/'Report.xlsx'),'running':False})
        url,error=self.backend.ready.get(timeout=20);self.assertIsNone(error)
        self.api=Api(url);self.backend.ready.put((url,None))
        self.browser=patch('webbrowser.open',side_effect=AssertionError('A native application must not open a browser'));self.browser.start()
        self.window=Window(self.backend,QSettings(str(self.root/'prefs.ini'),QSettings.IniFormat));self.window.show()
        wait_for(lambda:bool(self.window.state))
    def tearDown(self):
        self.window.close();APP.processEvents();self.browser.stop();self.tmp.cleanup()
    def refresh(self):
        self.window.signature=None;self.window.tick();wait_for(lambda:not self.window.polling)
    def add(self,**kwargs):
        return self.api.request('add-fault',dict(service_date='2026-10-16T12:00',worker_initials='BK',issue='Synthetic MLC mismatch',**kwargs))
    def test_native_form_edit_and_shared_conflict(self):
        def fill():
            dialog=self.window.findChild(QDialog)
            lines=dialog.findChildren(QLineEdit);lines[1].setText('AB');lines[2].setText('420219');lines[3].setText('MLC leaf')
            areas=dialog.findChildren(QPlainTextEdit);areas[0].setPlainText('Synthetic therapist fault');areas[1].setPlainText('Reset connection')
            dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.Save).click()
        QTimer.singleShot(100,fill);self.window.add_fault()
        self.refresh();rows=self.window.tab_tools['screenshots']['rows'];self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['work_performed'],'Reset connection')
        def edit():
            dialog=self.window.findChild(QDialog)
            # The closed Add fault dialog may remain as a QObject child until destruction.
            dialogs=[d for d in self.window.findChildren(QDialog) if d.isVisible()];dialog=dialogs[-1]
            dialog.findChild(QPlainTextEdit).setPlainText('Verified after reset')
            dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.Save).click()
        QTimer.singleShot(100,edit);self.window.edit_cell('screenshots',0,5);self.refresh()
        row=self.api.request('state')['records'][0];self.assertEqual(row['work_performed'],'Verified after reset')
        with self.assertRaises(ValueError):
            self.api.request('inline-update',{'id':row['id'],'field':'work_performed','value':'stale','expected':'Reset connection'})
        self.assertNotIn('service_status',[c[0] for c in FAULT_COLUMNS])
        self.assertNotIn('notes',[c[0] for c in FAULT_COLUMNS])
    def test_pdf_png_import_and_native_source_open(self):
        import pymupdf
        doc=pymupdf.open();page=doc.new_page();page.insert_text((72,72),'Error code: 420219');data=doc.tobytes();doc.close()
        result=self.api.request('upload',{'name':'synthetic-report.pdf','data':base64.b64encode(data).decode()});self.assertFalse(result['error'])
        png=(Path(__file__).parent/'assets'/'truebeam-icon.png').read_bytes()
        result=self.add(attachments=[{'name':'test.png','data':base64.b64encode(png).decode()}])
        self.refresh();fault=next(r for r in self.window.records if r['kind']=='screenshots');source=fault['sources'][0]
        self.assertEqual(self.api.request('source?id='+str(source['id']),raw=True),png)
        with patch('native_application.app_home',return_value=self.root/'client'),patch('native_application.QDesktopServices.openUrl',return_value=True) as opened:
            self.window.open_source(source);wait_for(lambda:opened.called)
            self.assertTrue(opened.call_args.args[0].isLocalFile())
        self.assertEqual(len(self.window.tab_tools['reports']['rows']),1)
    def test_sort_filter_dark_mode_and_download(self):
        self.add(fault_code='510102')
        self.api.request('add-fault',{'service_date':'2026-10-15','worker_initials':'CD','issue':'Other fault','fault_code':'420219'})
        self.refresh();tools=self.window.tab_tools['screenshots']
        self.window.sort('screenshots',1);self.assertEqual(tools['rows'][0]['fault_code'],'420219')
        tools['search'].setText('510102');self.assertEqual(len(tools['rows']),1)
        self.window.dark.setChecked(True);self.assertTrue(self.window.preferences.value('dark',type=bool))
        self.window.tabs.setCurrentIndex(1);APP.processEvents()
        self.assertLessEqual(sum(self.window.tables['screenshots'].columnWidth(i) for i in range(len(FAULT_COLUMNS))),self.window.tables['screenshots'].viewport().width()+1)
        path=self.root/'Faults.csv'
        with patch('native_application.QFileDialog.getSaveFileName',return_value=(str(path),'CSV')):self.window.download_table('screenshots')
        with path.open(encoding='utf-8-sig') as stream:rows=list(csv.reader(stream))
        self.assertIn('Resolution',rows[0]);self.assertNotIn('Status',rows[0]);self.assertEqual(len(rows),2)
    def test_export_snapshot_has_no_external_feed_dependency(self):
        self.add(fault_code='420219');records=self.api.request('state')['records']
        output=self.root/'Snapshot.xlsx';create_live_workbook(records,output,[],live=False)
        with zipfile.ZipFile(output) as z:
            self.assertNotIn('xl/connections.xml',z.namelist())
        from openpyxl import load_workbook
        wb=load_workbook(output);self.assertEqual(wb.sheetnames,['Service Reports','Fault Reports'])
        self.assertEqual(wb['Fault Reports']['B5'].value,'420219');wb.close()
if __name__=='__main__':unittest.main()
