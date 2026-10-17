"""Publish CSV data every minute; never replace an already-created, open Excel workbook."""
import csv, io, os, tempfile, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from report_export import write_report
MAIN='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/package/2006/relationships'
OFF='http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
def tag(name):return '{'+MAIN+'}'+name
def xml(element):
    ET.register_namespace('',element.tag[1:].split('}')[0])
    return ET.tostring(element,encoding='utf-8',xml_declaration=True)
SERVICE=[('service_date','Malfunction Start'),('machine_release','Machine Release'),('work_order','Work Order #'),('case_number','Case #'),('category','Category'),('issue','Subject'),('error_codes','Error Codes'),('work_performed','Summary of Work'),('parts_installed','Parts Installed'),('engineer','Field Engineer'),('work_hours','Total Work Hrs'),('downtime_hours','Agreed Downtime Hrs'),('service_status','Status'),('notes','Follow-up / Notes'),('sources','Source Report')]
FAULT=[('service_date','Event Date'),('error_codes','Error Codes'),('fault_name','Fault Name'),('severity','Severity'),('issue','Error Details'),('suggestions','SUGGESTED SOLUTION'),('work_performed','Resolution'),('worker_initials','Worker Initials'),('sources','Screenshot')]
def data_value(record,key):
    if key=='sources':return '\n'.join(s['path'] for s in record['sources'])
    if key=='error_codes':return '\n'.join(c.strip() for part in (record.get('fault_code',''),record.get('additional_codes','')) for c in str(part).replace(';',',').split(',') if c.strip())
    if key=='service_status':return record.get(key) or record.get('review','')
    return record.get(key,'') or ''
def atomic_write(path,data):
    fd,tmp=tempfile.mkstemp(dir=path.parent,suffix='.tmp')
    try:
        with os.fdopen(fd,'wb') as f:f.write(data)
        os.replace(tmp,path)
    finally:
        if os.path.exists(tmp):os.unlink(tmp)
def csv_data(records,columns):
    stream=io.StringIO(newline='');writer=csv.writer(stream)
    writer.writerow([name for _,name in columns])
    for record in records:
        # Text-query fields are explicitly imported as text; also guard ordinary CSV opens.
        row=[]
        for key,_ in columns:
            text=str(data_value(record,key))
            if text.startswith(('=','+','-','@','\t','\r')):text="'"+text
            row.append(text)
        writer.writerow(row)
    return ('\ufeff'+stream.getvalue()).encode('utf-8')
def create_live_workbook(records,output,feeds,live=True):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table,TableStyleInfo
    wb=Workbook();wb.remove(wb.active)
    for i,(kind,title,columns) in enumerate([('reports','Service Reports',SERVICE),('screenshots','Fault Reports',FAULT)],1):
        ws=wb.create_sheet(title);ws.sheet_view.showGridLines=False
        ws.append(['TrueBeam — '+title]);ws['A1'].font=Font(name='Arial',size=14,bold=True)
        ws.append(['Live report: data refreshes every minute. Make edits in the shared app. Enable external-data refresh in Excel if prompted.' if live else 'Exported snapshot of the current TrueBeam service and fault log.'])
        ws['A2'].font=Font(name='Arial',size=9,color='595959');ws.append([])
        ws.append([name for _,name in columns])
        rows=[r for r in records if r['kind']==kind]
        for r in rows:ws.append([data_value(r,key) for key,_ in columns])
        if not rows:ws.append(['']*len(columns))
        table=Table(displayName='TrueBeamLive'+str(i),ref='A4:'+get_column_letter(len(columns))+str(ws.max_row))
        if live:table.tableType='queryTable';table.connectionId=i
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        table._initialise_columns()
        for index,(_,name) in enumerate(columns):
            table.tableColumns[index].name=name
            if live:table.tableColumns[index].queryTableFieldId=index+1
        ws.add_table(table);ws.freeze_panes='A5'
        for c in ws[4]:
            c.fill=PatternFill('solid',fgColor='1F4E78');c.font=Font(name='Arial',bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True)
        for row in ws.iter_rows(min_row=5):
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.data_type='s'
        for c,(key,_) in enumerate(columns,1):ws.column_dimensions[get_column_letter(c)].width=60 if key in ('issue','suggestions','work_performed') else 22
    fd,tmp=tempfile.mkstemp(dir=output.parent,suffix='.xlsx');os.close(fd)
    try:
        wb.save(tmp)
        if not live:
            os.replace(tmp,output);return
        with zipfile.ZipFile(tmp) as z:parts={n:z.read(n) for n in z.namelist()}
        style=ET.fromstring(parts['xl/styles.xml'])
        font_order=['b','i','strike','condense','extend','outline','shadow','u','vertAlign','sz','color','name','family','charset','scheme']
        for font in style.find(tag('fonts')):
            children=list(font);children.sort(key=lambda e:font_order.index(e.tag.split('}')[-1]) if e.tag.split('}')[-1] in font_order else 99)
            font[:]=children
        parts['xl/styles.xml']=xml(style)
        connections=ET.Element(tag('connections'))
        for i,(columns,feed) in enumerate(zip((SERVICE,FAULT),feeds),1):
            conn=ET.SubElement(connections,tag('connection'),id=str(i),name='TrueBeam data '+str(i),type='6',refreshedVersion='6',minRefreshableVersion='3',background='1',saveData='1',interval='1',refreshOnLoad='1')
            text=ET.SubElement(conn,tag('textPr'),prompt='0',characterSet='UTF-8',sourceFile=str(feed),firstRow='1',delimited='1',comma='1',tab='0',qualifier='doubleQuote')
            fields=ET.SubElement(text,tag('textFields'),count=str(len(columns)))
            for n in range(len(columns)):ET.SubElement(fields,tag('textField'),type='text',position=str(n))
            query=ET.Element(tag('queryTable'),name='TrueBeamLive'+str(i),connectionId=str(i),backgroundRefresh='1',refreshOnLoad='1',adjustColumnWidth='0',preserveFormatting='1')
            refresh=ET.SubElement(query,tag('queryTableRefresh'),nextId=str(len(columns)+1))
            qfields=ET.SubElement(refresh,tag('queryTableFields'),count=str(len(columns)))
            for n,(_,name) in enumerate(columns,1):ET.SubElement(qfields,tag('queryTableField'),id=str(n),name=name,tableColumnId=str(n))
            parts['xl/queryTables/queryTable'+str(i)+'.xml']=xml(query)
            rels=ET.Element('{'+REL+'}Relationships')
            ET.SubElement(rels,'{'+REL+'}Relationship',Id='rIdQuery',Type=OFF+'queryTable',Target='../queryTables/queryTable'+str(i)+'.xml')
            parts['xl/tables/_rels/table'+str(i)+'.xml.rels']=xml(rels)
        parts['xl/connections.xml']=xml(connections)
        rels=ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
        ET.SubElement(rels,'{'+REL+'}Relationship',Id='rIdConnections',Type=OFF+'connections',Target='connections.xml')
        parts['xl/_rels/workbook.xml.rels']=xml(rels)
        types=ET.fromstring(parts['[Content_Types].xml'])
        for name,kind in [('/xl/connections.xml','connections'),('/xl/queryTables/queryTable1.xml','queryTable'),('/xl/queryTables/queryTable2.xml','queryTable')]:
            ET.SubElement(types,'{'+CT+'}Override',PartName=name,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.'+kind+'+xml')
        parts['[Content_Types].xml']=xml(types)
        with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as z:
            for name,data in parts.items():z.writestr(name,data)
        # Initial creation only; subsequent publications never modify this file.
        if output.exists():raise FileExistsError('Choose a new report filename; the existing workbook is preserved.')
        os.replace(tmp,output)
    finally:
        if os.path.exists(tmp):os.unlink(tmp)
def publish_report(records,output):
    output=Path(output)
    if not output.parent.is_dir():raise ValueError('Report folder unavailable: '+str(output.parent))
    folder=output.parent/(output.stem+'_data');folder.mkdir(exist_ok=True)
    feeds=[folder/'ServiceReports.csv',folder/'FaultReports.csv']
    for kind,columns,path in zip(('reports','screenshots'),(SERVICE,FAULT),feeds):
        atomic_write(path,csv_data([r for r in records if r['kind']==kind],columns))
    if not output.exists():create_live_workbook(records,output,feeds)
    return output
