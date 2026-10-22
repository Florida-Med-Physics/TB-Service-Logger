import csv,json,os,sqlite3,tempfile,threading,unittest
from datetime import datetime
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.request import Request,urlopen
from urllib.error import HTTPError
from application_host import SharedService,Server,Conflict,DEFAULTS
from live_report import publish_report
class SharedTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
  for name in ('reports','screens','network'):(self.root/name).mkdir()
  self.defaults={'reports':str(self.root/'reports'),'screenshots':str(self.root/'screens'),'output':str(self.root/'network'/'Report.xlsx'),'running':False}
  self.service=SharedService(self.root/'home',defaults=self.defaults,shared_dir=self.root/'network')
  self.server=Server(('127.0.0.1',0),self.service);self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start()
  self.url='http://127.0.0.1:'+str(self.server.server_port)
 def tearDown(self):
  self.server.shutdown();self.server.server_close();self.service.close();self.tmp.cleanup()
 def call(self,route,body=None,token=True):
  return urlopen(Request(self.url+'/api/'+route,data=None if body is None else json.dumps(body).encode(),headers={'X-Local-Token':self.service.token if token else '',**({'Content-Type':'application/json'} if body is not None else {})}),timeout=15)
 def fault(self):
  return json.load(self.call('add-fault',{'service_date':'2026-10-16T12:00','worker_initials':'BK','issue':'MLC fault'}))['id']
 def test_evening_email_settings_persist_and_survive_folder_changes(self):
  recipients=['physics@example.org','therapy@example.org']
  self.assertTrue(json.load(self.call('email-settings',{'enabled':True,'recipients':recipients+['physics@example.org']}))['ok'])
  self.call('configure',{'reports':self.defaults['reports'],'screenshots':self.defaults['screenshots']}).read()
  config=json.load(self.call('state'))['config']
  self.assertTrue(config['evening_email']);self.assertEqual(config['email_recipients'],recipients)
  self.assertEqual(json.loads((self.root/'network'/'settings.json').read_text())['email_recipients'],recipients)
  for body in ({'enabled':True,'recipients':[]},{'enabled':False,'recipients':['not-an-address']}):
   with self.assertRaises(HTTPError) as error:self.call('email-settings',body)
   self.assertEqual(error.exception.code,400)
 def test_evening_email_sends_once_after_time_without_screenshot_text(self):
  sent=[];fail=[]
  class FakeSMTP:
   def __init__(self,host,port,timeout):self.address=(host,port)
   def __enter__(self):return self
   def __exit__(self,*args):pass
   def ehlo(self):pass
   def has_extn(self,name):return False
   def send_message(self,message):
    if fail:raise OSError('relay refused connection')
    sent.append((self.address,message))
  self.service.environ={'TRUEBEAM_SMTP_HOST':'relay.example.org','TRUEBEAM_EMAIL_FROM':'truebeam@example.org'}
  self.fault();self.call('email-settings',{'enabled':True,'recipients':['physics@example.org'],'time':'18:00'}).read()
  today=datetime.now().replace(second=0,microsecond=0)
  with patch('smtplib.SMTP',FakeSMTP):
   self.service.now=lambda:today.replace(hour=17,minute=59);self.service.email_tick();self.assertEqual(sent,[])
   fail.append(True);self.service.now=lambda:today.replace(hour=18,minute=1);self.service.email_tick()
   self.assertIn('relay refused',json.load(self.call('state'))['email_status']['last_error'])
   fail.clear();self.service.email_tick();self.assertEqual(sent,[])
   self.service.now=lambda:today.replace(hour=18,minute=17);self.service.email_tick();self.service.email_tick()
  self.assertEqual(len(sent),1);address,message=sent[0]
  self.assertEqual(address,('relay.example.org',25));self.assertEqual(message['To'],'physics@example.org')
  body=message.get_body(('plain',)).get_content()
  self.assertEqual(message['Subject'],'UF TB MACHINE LOG: '+today.strftime('%m/%d/%Y'))
  self.assertIn('1 fault',body);self.assertIn('BK',body);self.assertIn('Error Details: MLC fault',body);self.assertIn('Past 5 faults (1)',body)
  status=json.load(self.call('state'))['email_status'];self.assertEqual(status['last_error'],'');self.assertTrue(status['last_sent_at'])
 def test_test_email_explains_missing_mail_server(self):
  self.service.environ={};self.service.outlook_check=lambda:False
  self.call('email-settings',{'enabled':True,'recipients':['physics@example.org']}).read()
  self.assertFalse(json.load(self.call('state'))['email_status']['configured'])
  with self.assertRaises(HTTPError) as error:self.call('email-test',{})
  self.assertIn('TRUEBEAM_SMTP_HOST',json.load(error.exception)['error'])
 def test_test_email_uses_outlook_without_mail_server(self):
  runs=[]
  def fake_run(command,**options):
   runs.append(json.loads(Path(command[-1]).read_text(encoding='utf-8')))
   return type('Result',(),{'returncode':0,'stderr':''})()
  self.service.environ={};self.service.outlook_check=lambda:True;self.fault()
  self.call('email-settings',{'enabled':True,'recipients':['physics@example.org','therapy@example.org']}).read()
  self.assertEqual(json.load(self.call('state'))['email_status']['via'],'Outlook on this host')
  with patch('evening_email.subprocess.run',fake_run):self.assertEqual(json.load(self.call('email-test',{}))['entries'],1)
  self.assertEqual(runs[0]['to'],'physics@example.org; therapy@example.org')
  self.assertTrue(runs[0]['subject'].startswith('TEST - UF TB MACHINE LOG: '));self.assertIn('MLC fault',runs[0]['html'])
 def test_email_tables_list_past_five_faults_and_withhold_patient_text(self):
  import evening_email
  faults=[dict(id=i,kind='screenshots',service_date='2026-10-%02dT10:00:00'%i,fault_code=str(420200+i),issue='Interlock %d'%i,sources=[{'id':i}]) for i in range(1,8)]
  faults[6]['issue']='Patient: SAMPLE, SYNTHETIC  MRN 0000000'
  report=dict(id=9,kind='reports',service_date='2026-10-20',work_order='WO-1',sources=[])
  now=datetime(2026,10,21,18,0)
  message=evening_email.build_message([faults[6],report],faults+[report],['physics@example.org'],'',datetime(2026,10,20,18,0),now)
  body=message.get_body(('plain',)).get_content();recent=body.split('Past 5 faults (5)')[1]
  self.assertEqual(message['Subject'],'UF TB MACHINE LOG: 10/21/2026')
  self.assertEqual([c for c in ('420207','420206','420205','420204','420203','420202') if c in recent],['420207','420206','420205','420204','420203'])
  self.assertNotIn('SYNTHETIC',body);self.assertIn(evening_email.WITHHELD,body)
  self.assertIn('Event time: 10/07/2026 10:00',body);self.assertIn('Work order: WO-1',body);self.assertIn('Source count: 1',body)
 def test_fault_status_starts_open_and_saves_resolved_in_email(self):
  import evening_email
  rid=self.fault()
  self.assertTrue(json.load(self.call('inline-update',{'id':rid,'field':'service_status','value':'Resolved','expected':''}))['ok'])
  self.assertTrue(json.load(self.call('inline-update',{'id':rid,'field':'current_status','value':'Waiting on field engineer','expected':''}))['ok'])
  records=self.service.state()['records'];self.assertEqual(records[0]['service_status'],'Resolved')
  other=dict(id=99,kind='screenshots',service_date='2026-10-15T10:00:00',issue='Other fault',sources=[])
  message=evening_email.build_message(records+[other],records+[other],['physics@example.org'],'',datetime(2026,10,15,18,0),datetime(2026,10,16,18,0))
  body=message.get_body(('plain',)).get_content();page=message.get_body(('html',)).get_content()
  self.assertIn('Status: Resolved',body);self.assertIn('Status: Open',body)
  self.assertIn('Varian Suggested Solution',body);self.assertIn('Final Resolution',body);self.assertIn('Current Status: Waiting on field engineer',body);self.assertIn('Initials: BK',body)
  self.assertIn('background:#CDEED6',page);self.assertIn('background:#FDE2C4',page)
 def test_downtime_alert_recipients_default_list(self):
  from application_host import DEFAULT_ALERT_RECIPIENTS
  self.assertEqual(json.load(self.call('state'))['config']['alert_recipients'],['Physics@floridaproton.org','tfleming@floridaproton.org','Therapists@floridaproton.org'])
  self.assertEqual(DEFAULT_ALERT_RECIPIENTS,json.loads((self.root/'network'/'settings.json').read_text())['alert_recipients'])
 def test_downtime_alert_sends_scripted_email_and_refuses_identifiers(self):
  runs=[]
  def fake_run(command,**options):
   runs.append(json.loads(Path(command[-1]).read_text(encoding='utf-8')))
   return type('Result',(),{'returncode':0,'stderr':''})()
  self.service.environ={};self.service.outlook_check=lambda:True
  self.call('email-settings',{'enabled':False,'recipients':['evening@example.org']}).read()
  self.call('alert-settings',{'recipients':['physics@example.org','therapy@example.org','physics@example.org']}).read()
  self.call('configure',{'reports':self.defaults['reports'],'screenshots':self.defaults['screenshots']}).read()
  config=json.load(self.call('state'))['config']
  self.assertEqual(config['alert_recipients'],['physics@example.org','therapy@example.org']);self.assertEqual(config['email_recipients'],['evening@example.org'])
  with self.assertRaises(HTTPError):self.call('alert-settings',{'recipients':['not-an-address']})
  alert={'site':'Truebeam','status':'Machine down','started':'2026-09-12T10:00','issue':'Synthetic MLC trajectory error','details':'Synthetic fault after QA.',
         'vendor':'Varian helpline called, case # 00000000','patient_impact':'2 patients moved to another machine','plan':'Field engineer contacted','physics_impact':'No PSQA until back up','reported_by':'bk'}
  with patch('evening_email.subprocess.run',fake_run):
   self.assertEqual(json.load(self.call('downtime-alert',alert))['recipients'],2)
   with self.assertRaises(HTTPError) as error:self.call('downtime-alert',{**alert,'patient_impact':'MRN 0000000 moved'})
   self.service.now=lambda:datetime(2026,9,12,14,30)
   self.call('downtime-alert',{'status':'Machine down','recipients':'oncall@example.org; physics@example.org'}).read()
   with self.assertRaises(HTTPError) as bad:self.call('downtime-alert',{'status':'Machine down','recipients':'not-an-address'})
  self.assertIn('Patient Impact',json.load(error.exception)['error']);self.assertIn('not-an-address',json.load(bad.exception)['error']);self.assertEqual(len(runs),2)
  self.assertEqual(runs[0]['subject'],'DOWNTIME ALERT: Truebeam 9/12 @10am');self.assertEqual(runs[0]['to'],'physics@example.org; therapy@example.org')
  self.assertEqual(runs[1]['subject'],'DOWNTIME ALERT: Truebeam 9/12 @2:30pm');self.assertEqual(runs[1]['to'],'oncall@example.org; physics@example.org')
  for text in ('Vendor Actions to Date','Patient Impact','Plan Going Forward','Physics Impact','09/12/2026 10:00','BK'):self.assertIn(text,runs[0]['html'])
 def test_export_report_downloads_current_log(self):
  from io import BytesIO
  from openpyxl import load_workbook
  self.fault()
  wb=load_workbook(BytesIO(self.call('excel').read()));ws=wb['Fault Reports']
  self.assertEqual(ws.cell(5,[c.value for c in ws[4]].index('Error Details')+1).value,'MLC fault')
 def test_two_clients_merge_different_fields_and_reject_same_field_conflict(self):
  rid=self.fault()
  with ThreadPoolExecutor() as pool:
   results=list(pool.map(lambda p:json.load(self.call('inline-update',p)),[
    {'id':rid,'field':'work_performed','value':'Reset connection','expected':''},
    {'id':rid,'field':'fault_code','value':'420219','expected':''}]))
  self.assertTrue(all(r['ok'] for r in results))
  with self.assertRaises(HTTPError) as error:self.call('inline-update',{'id':rid,'field':'work_performed','value':'Old edit','expected':''})
  self.assertEqual(error.exception.code,409)
  r=json.load(self.call('state'))['records'][0];self.assertEqual(r['work_performed'],'Reset connection');self.assertEqual(r['fault_code'],'420219')
 def test_shared_backup_and_paths_survive_handoff(self):
  rid=self.fault();backup=self.root/'network'/'shared-data.sqlite3';self.assertTrue(backup.exists())
  with sqlite3.connect(backup) as db:self.assertEqual(db.execute('SELECT count(*) FROM records').fetchone()[0],1)
  # Close the inspection handle on Windows.
  db.close()
  self.assertEqual(json.loads((self.root/'network'/'settings.json').read_text())['reports'],self.defaults['reports'])
 def test_report_publication_never_writes_open_workbook(self):
  self.fault();output=Path(self.defaults['output']);publish_report(self.service.state()['records'],output);before=output.read_bytes();stamp=output.stat().st_mtime_ns
  rid=self.service.state()['records'][0]['id'];self.service.update({'id':rid,'field':'work_performed','value':'New fix','expected':''})
  with open(output,'rb') as held:
   publish_report(self.service.state()['records'],output)
   self.assertEqual(held.read(),before)
  self.assertEqual(output.stat().st_mtime_ns,stamp)
  feed=output.parent/(output.stem+'_data')/'FaultReports.csv'
  self.assertIn('New fix',feed.read_text(encoding='utf-8-sig'))
 def test_reuse_existing_report_after_restart_and_different_saved_path(self):
  self.fault();output=Path(self.defaults['output']);publish_report(self.service.state()['records'],output)
  before=output.read_bytes();stamp=output.stat().st_mtime_ns
  # A different profile/previous setting must not reject the existing shared report.
  self.service.engine.config['output']=str(self.root/'old-report.xlsx')
  with output.open('rb') as held:
   response=json.load(self.call('configure',self.defaults))
   self.assertTrue(response['ok']);self.service.tick()
   self.assertEqual(held.read(),before)
  self.assertEqual(output.stat().st_mtime_ns,stamp)
  self.service.close()
  self.service=SharedService(self.root/'home',defaults=self.defaults,shared_dir=self.root/'network')
  self.server.service=self.service
  self.assertEqual(self.service.engine.config['output'],str(output.resolve()))
  self.assertTrue(json.load(self.call('configure',self.defaults))['ok'])
  self.assertEqual(output.read_bytes(),before)
  feed=output.parent/(output.stem+'_data')/'FaultReports.csv'
  self.assertTrue(feed.exists())
 def test_reusing_unrelated_existing_workbook_does_not_replace_it(self):
  output=self.root/'network'/'Existing.xlsx';output.write_bytes(b'existing workbook bytes')
  config=dict(self.defaults,output=str(output))
  self.assertTrue(json.load(self.call('configure',config))['ok'])
  self.service.tick();self.assertEqual(output.read_bytes(),b'existing workbook bytes')
 def test_filter_existing_imports_and_use_screenshot_filename_time(self):
  name='Print Screen 2026_10_02 08_37_27.png'
  documents={
   'PPG-HT-G.pdf':'Product guide PPG-HT-G',
   'WO-004612326.pdf':'Work Order Information\nWork Order WO-004612326\nSubject Artifacts on conebeam\nClosure Summary Inspected imaging\nWork Order Times',
   'normal.png':'Treatment screen - no active faults',
   name:'Id: 136011\nName: SPV.SW.BEAM.MotionAxisNotAtPlan\nDescription: Treatment axis is not at plan\nDateTime: 29-Sep-2026 02:36:30 PM\nAdditional fault codes: 136011'}
  for filename in documents:
   folder=self.root/('reports' if filename.endswith('.pdf') else 'screens')
   path=folder/filename;path.write_text('synthetic '+filename);os.utime(path,(1,1))
  self.service.engine.configure(self.defaults,reuse_existing=True)
  with patch('core.extract',side_effect=lambda path:documents[path.name]):
   self.service.engine.scan(auto_export=False);self.service.engine.scan(auto_export=False)
  self.service.engine.config['running']=False
  # Simulate an old database row; the next state must derive its date from the file.
  self.service.engine.db.execute("UPDATE records SET service_date='wrong OCR date' WHERE kind='screenshots'")
  self.service.engine.db.commit()
  state=self.service.state()
  self.assertEqual(len(state['records']),3);self.assertEqual(len(state['excluded']),1)
  fault=next(r for r in state['records'] if r['kind']=='screenshots' and r['fault_code'])
  self.assertEqual(fault['service_date'],'2026-10-02T08:37:27')
  self.assertEqual(len(self.service.engine.rows()),4)
  publish_report(self.service.state()['records'],Path(self.defaults['output']))
  feed=self.root/'network'/'Report_data'/'ServiceReports.csv'
  self.assertNotIn('PPG-HT-G',feed.read_text(encoding='utf-8-sig'))
  self.assertIn('WO-004612326',feed.read_text(encoding='utf-8-sig'))
 def test_authentication_and_cross_origin(self):
  with self.assertRaises(HTTPError) as error:self.call('state',token=False)
  self.assertEqual(error.exception.code,403)
  request=Request(self.url+'/api/add-fault',data=b'{}',headers={'X-Local-Token':self.service.token,'Origin':'https://example.com','Content-Type':'application/json'})
  with self.assertRaises(HTTPError) as error:urlopen(request)
  self.assertEqual(error.exception.code,403)
 def test_tenure_defaults(self):
  self.assertEqual(DEFAULTS['reports'],r'U:\QA\SystemQA\LINAC QA\TrueBeam\Service reports')
  self.assertEqual(DEFAULTS['screenshots'],r'W:\Varian\Transfer\TDS\H192905\PrintScreen')
if __name__=='__main__':unittest.main()
